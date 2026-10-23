"""Flag shapes outside the slide or overlapping the footer strip. Usage: check_pptx.py <file.pptx>"""
import sys
from pptx import Presentation
from pptx.util import Emu
prs = Presentation(sys.argv[1]); W, H = prs.slide_width, prs.slide_height
bad = 0
for n, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.left is None: continue
        r, b = sh.left + sh.width, sh.top + sh.height
        footer_hit = (s.slide_layout.name == "Title Only" and b > Emu(int(6.45 * 914400)) and sh.shape_type != 14)
        if sh.left < -9525 or sh.top < -9525 or r > W + 9525 or b > H + 9525 or footer_hit:
            bad += 1; print(f"slide {n}: {sh.name} {Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f} -> {Emu(r).inches:.2f},{Emu(b).inches:.2f}")
print("issues:", bad)
