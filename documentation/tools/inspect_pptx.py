"""Print layouts, placeholders and slides of a pptx. Usage: inspect_pptx.py <file.pptx>"""
import sys
from pptx import Presentation
from pptx.util import Emu

prs = Presentation(sys.argv[1])
print("size", Emu(prs.slide_width).inches, Emu(prs.slide_height).inches)
for i, lay in enumerate(prs.slide_layouts):
    print(f"LAYOUT {i}: {lay.name}")
    for sh in lay.placeholders:
        print(f"   ph idx={sh.placeholder_format.idx} type={sh.placeholder_format.type} {sh.name} "
              f"{Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f} {Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f}")
for n, s in enumerate(prs.slides, 1):
    print(f"SLIDE {n}: layout={s.slide_layout.name}")
    for sh in s.shapes:
        t = sh.text_frame.text[:60].replace("\n", " | ") if sh.has_text_frame else ""
        print(f"   {sh.shape_type} {sh.name} ph={sh.is_placeholder} {Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f} {Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f} '{t}'")
