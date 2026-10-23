"""Build presentation/PV_Digital_Twin.pptx from the DUT template, Mermaid diagrams and Wireloom wireframes.

Usage: build_presentation.py <template.pptx>
Block diagrams and the SDM circuit: presentation/typst/*.typ (Typst + CeTZ) -> presentation/png/*.png
Flow, state and sequence diagrams: presentation/diagrams/*.mmd (Mermaid) -> presentation/png/*.png (via mmdc)
Wireloom sources: wireframes/src/*.wireloom   -> wireframes/svg/*.svg (build_wireframes.py) -> png
"""
import subprocess
import sys
from pathlib import Path

import cairosvg
from PIL import Image
from pptx import Presentation
from pptx.opc.packuri import PackURI
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import pvgraphics  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
TOOLS = ROOT / "tools"
DIAGRAMS = ROOT / "presentation" / "diagrams"
PNG = ROOT / "presentation" / "png"
WF_SVG = ROOT / "wireframes" / "svg"
OUT = ROOT / "presentation" / "PV_Digital_Twin.pptx"

GREEN = RGBColor(0x13, 0x7F, 0x36)
GOLD = RGBColor(0xD7, 0xB6, 0x80)
BLACK = RGBColor(0, 0, 0)
GREY = RGBColor(0x55, 0x55, 0x55)

# Content area between the title and the template's footer strip (y = 6.45 in).
BOX_L, BOX_T, BOX_W, BOX_H = 0.92, 1.45, 11.5, 4.35


def render_diagram(name: str) -> Path:
    """Block diagrams are Typst + CeTZ when a .typ exists, otherwise Mermaid."""
    if (ROOT / "presentation" / "typst" / f"{name}.typ").exists():
        return pvgraphics.typst_png(name)
    return render_mermaid(name)


def render_mermaid(name: str) -> Path:
    src, out = DIAGRAMS / f"{name}.mmd", PNG / f"{name}.png"
    if not out.exists() or out.stat().st_mtime < src.stat().st_mtime:
        PNG.mkdir(exist_ok=True)
        subprocess.run(
            [str(TOOLS / "node_modules/.bin/mmdc"), "-i", str(src), "-o", str(out),
             "-c", str(TOOLS / "mermaid-config.json"), "-p", str(TOOLS / "puppeteer.json"),
             "-s", "3", "-b", "white", "-q"],
            check=True)
    return out


def render_wireframe(name: str) -> Path:
    src, out = WF_SVG / f"{name}.svg", PNG / f"wf_{name}.png"
    if not out.exists() or out.stat().st_mtime < src.stat().st_mtime:
        PNG.mkdir(exist_ok=True)
        cairosvg.svg2png(url=str(src), write_to=str(out), output_width=2200, background_color="white")
    return out


def fit_picture(slide, path, left, top, width, height):
    iw, ih = Image.open(path).size
    scale = min(width / iw, height / ih)
    w, h = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), Inches(left + (width - w) / 2), Inches(top + (height - h) / 2),
                                    Inches(w), Inches(h))


def add_text(slide, left, top, width, height, paras, size=18, color=BLACK, bold=False, italic=False,
             align=PP_ALIGN.LEFT, bullets=False):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, text in enumerate(paras if isinstance(paras, list) else [paras]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = ("•  " + text) if bullets else text
        run.font.size, run.font.bold, run.font.italic = Pt(size), bold, italic
        run.font.color.rgb = color
    return tb


def set_title(slide, text):
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.92), Inches(0.4), Inches(11.5), Inches(0.9)
    t.text_frame.text = text
    for p in t.text_frame.paragraphs:
        for r in p.runs:
            r.font.size, r.font.bold = Pt(32), True
            r.font.color.rgb = GREEN


def notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


class Deck:
    def __init__(self, template):
        self.prs = Presentation(template)
        s = self.prs.slides
        self.cover, self.thanks = s[0], s[7]
        self.section_imgs = [s[i].placeholders[13].image.blob for i in (1, 3, 5)]
        ids = self.prs.slides._sldIdLst
        for sld in list(ids)[1:7]:           # drop template section/content samples
            self.prs.part.drop_rel(sld.rId)
            ids.remove(sld)
        # new slides are numbered count+1, which would collide with the kept 'Thank you' part
        self.thanks.part.partname = PackURI('/ppt/slides/slide999.xml')
        self.layout = {l.name: l for l in self.prs.slide_layouts}
        self.section_no = 0

    def section(self, title):
        s = self.prs.slides.add_slide(self.layout["Section Header"])
        s.shapes.title.text = title
        blob = self.section_imgs[self.section_no % len(self.section_imgs)]
        tmp = PNG / "_section.jpg"
        tmp.write_bytes(blob)
        s.placeholders[13].insert_picture(str(tmp))
        self.section_no += 1
        s.placeholders[14].text = f"{self.section_no:02d}"
        return s

    def bullets(self, title, items, note=""):
        s = self.prs.slides.add_slide(self.layout["Title Only"])
        set_title(s, title)
        add_text(s, 0.92, 1.5, 11.5, 4.7, items, size=22, bullets=True)
        notes(s, note)
        return s

    def diagram(self, title, name, caption="", note=""):
        s = self.prs.slides.add_slide(self.layout["Title Only"])
        set_title(s, title)
        fit_picture(s, render_diagram(name), BOX_L, BOX_T, BOX_W, BOX_H if caption else BOX_H + 0.5)
        if caption:
            add_text(s, BOX_L, 5.85, BOX_W, 0.5, caption, size=18, color=GREY, italic=True, align=PP_ALIGN.CENTER)
        notes(s, note)
        return s

    def diagram_side(self, title, name, items, note="", heading=None):
        s = self.prs.slides.add_slide(self.layout["Title Only"])
        set_title(s, title)
        fit_picture(s, render_diagram(name), 0.6, BOX_T, 6.4, 4.85)
        self._side_text(s, items, heading)
        notes(s, note)
        return s

    def wireframe(self, title, name, items, note=""):
        s = self.prs.slides.add_slide(self.layout["Title Only"])
        set_title(s, title)
        fit_picture(s, render_wireframe(name), 0.5, BOX_T, 7.9, 4.9)
        self._side_text(s, items, "What to point out")
        notes(s, note)
        return s

    def _side_text(self, slide, items, heading):
        y = 1.6
        if heading:
            add_text(slide, 8.6, y, 4.3, 0.5, heading, size=20, bold=True, color=GREEN)
            y += 0.55
        add_text(slide, 8.6, y, 4.3, 6.3 - y, items, size=17, bullets=True)

    def equations(self, title, rows, note=""):
        """rows: (caption, key, latex). Equations are typeset images, not text."""
        s = self.prs.slides.add_slide(self.layout["Title Only"])
        set_title(s, title)
        row_h = 4.7 / len(rows)
        for i, (caption, key, latex) in enumerate(rows):
            top = 1.5 + i * row_h
            add_text(s, 0.92, top + row_h / 2 - 0.3, 3.2, 0.6, caption, size=18, bold=True, color=GREEN)
            fit_picture(s, pvgraphics.equation_png(key, latex), 4.3, top + 0.05, 8.1, row_h - 0.2)
        notes(s, note)
        return s

    def picture(self, title, path, eq_key=None, note=""):
        """Large picture, optionally with a typeset equation underneath."""
        s = self.prs.slides.add_slide(self.layout["Title Only"])
        set_title(s, title)
        fit_picture(s, path, BOX_L, BOX_T, BOX_W, 3.6 if eq_key else BOX_H + 0.5)
        if eq_key:
            fit_picture(s, pvgraphics.equation_png(eq_key), BOX_L, 5.15, BOX_W, 1.1)
        notes(s, note)
        return s

    def table(self, title, header, rows, note=""):
        s = self.prs.slides.add_slide(self.layout["Title Only"])
        set_title(s, title)
        shape = s.shapes.add_table(len(rows) + 1, len(header), Inches(0.92), Inches(1.5), Inches(11.5),
                                   Inches(0.6 * (len(rows) + 1)))
        tbl = shape.table
        for c, h in enumerate(header):
            self._cell(tbl.cell(0, c), h, True, GREEN, RGBColor(255, 255, 255))
        for r, row in enumerate(rows, 1):
            for c, v in enumerate(row):
                self._cell(tbl.cell(r, c), v, c == 0, RGBColor(0xF7, 0xF3, 0xEA) if r % 2 else RGBColor(255, 255, 255), BLACK)
        notes(s, note)
        return s

    @staticmethod
    def _cell(cell, text, bold, fill, color):
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
        cell.text = text
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size, r.font.bold = Pt(18), bold
                r.font.color.rgb = color

    def finish(self):
        ids = self.prs.slides._sldIdLst
        thanks = [e for e in ids if e.rId == self._rid(self.thanks)][0]
        ids.remove(thanks)
        ids.append(thanks)
        self.prs.save(OUT)

    def _rid(self, slide):
        for rel in self.prs.part.rels.values():
            if rel.reltype.endswith("/slide") and rel.target_part is slide.part:
                return rel.rId


def main(template):
    d = Deck(template)

    # Cover (template slide 1)
    ph = {p.placeholder_format.idx: p for p in d.cover.placeholders}
    ph[0].text = "Physics-Informed Digital Twin for PV State Estimation"
    ph[1].text = "Low-cost PV monitoring with the Single Diode Model"
    for sh in d.cover.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == "Presenter":
            sh.text_frame.text = "Department of Electronic & Computer Engineering"
            sh.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    for p in ph[0].text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(26)
    for p in ph[1].text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(16)

    # 01 The project
    d.section("The Project")
    d.bullets("Problem and objective", [
        "PV output drifts with irradiance and temperature, so a low reading alone does not mean a fault.",
        "A physics model gives the expected power for the measured conditions.",
        "Comparing measured and expected power shows when the system is abnormal.",
        "Goal: a low-cost digital twin that supports early identification of abnormal operation.",
    ], "Frame the problem: raw power is not enough, you need an expectation under the same conditions.")
    d.diagram("The twin at a glance", "d01-system-glance",
              "Same conditions go to the plant and the twin. The gap between them is the signal.",
              "Read left to right. Both the plant and the twin see the same G and T.")
    d.diagram("Planned HIL hardware setup", "d02-hardware",
              "Thick arrows: power path. Thin arrows: measurement and data signals.",
              "The intended hardware-in-the-loop setup: a real PV module and sensors feed the twin through the ESP32. "
              "Not assembled yet. The load type and the Wi-Fi link are assumptions to confirm.")
    d.diagram("Why simulation now", "d03-sim-to-hw",
              "A data interface means the hardware swap changes one box and nothing else.",
              "Hardware is unavailable, so a virtual plant feeds the same interface the ESP32 will use later.")

    # 02 The digital twin
    d.section("The Digital Twin")
    d.picture("Single Diode Model circuit", pvgraphics.circuit_png(), eq_key="kcl",
              note="The photocurrent splits between the diode, the shunt path and the output. Equivalent circuit of the module. Rs is the series loss, Rsh the shunt loss, D the diode.")
    d.equations("Model equations", [
        ("Single Diode Model", "sdm", r"I = I_{ph} - I_0\left(e^{\frac{V + I R_s}{n N_s V_t}} - 1\right) - \frac{V + I R_s}{R_{sh}}"),
        ("Photocurrent", "iph", r"I_{ph} = \frac{G}{G_{ref}}\left[I_{sc} + K_i\,(T - T_{ref})\right]"),
        ("Thermal voltage", "vt", r"V_t = \frac{k\,T}{q}"),
        ("Power deviation", "dev", r"\Delta = \frac{P_{meas} - P_{pred}}{P_{pred}} \times 100\%"),
    ], "The SDM is solved for I at each sample. Iph follows irradiance and temperature.")
    d.diagram("From conditions to a predicted curve", "d05-calc-curve",
              "Datasheet values fix the parameters. G and T move the curve.",
              "Chunk 1 of 2: inputs to I-V and P-V curve.")
    d.diagram("From prediction to state", "d06-calc-state",
              "Deviation is the percentage gap between measured and predicted power.",
              "Chunk 2 of 2: prediction to deviation to state.")
    d.diagram_side("Three operating states", "d07-states", [
        "Normal: deviation within the warning limit",
        "Warning: above the warning limit",
        "Abnormal: above the abnormal limit",
        "Limits come from the healthy-run error distribution, not hand-picked.",
    ], "State estimator.", "State logic")
    d.diagram("Plant and twin are separate models", "d08-plant-vs-twin",
              "If both were identical and noise-free, every error metric would be zero and prove nothing.",
              "The plant has perturbed parameters, faults and sensor noise. The twin uses nominal parameters.")
    d.table("Faults injected for testing", ["Fault", "Effect on the plant"], [
        ["Partial shading", "Lower photocurrent on part of the module"],
        ["Soiling", "Gradual loss of effective irradiance"],
        ["Series resistance rise", "Lower fill factor, less power"],
        ["Bypass diode failure", "Voltage step and power drop"],
        ["Sensor drift", "Biased measurement, false deviation"],
    ], "Each fault is labelled, so detection can be scored.")
    d.diagram("Validation: healthy data first", "d10-validation-healthy",
              "Error metrics and thresholds come from the same healthy runs.",
              "Chunk 1 of 2.")
    d.diagram("Validation: then faults", "d11-validation-faults",
              "Detection is scored on labelled faults, and thresholds are re-tuned if needed.",
              "Chunk 2 of 2.")

    d.equations("Validation metrics", [
        ("Mean absolute error", "mae", r"\mathrm{MAE} = \frac{1}{N}\sum_{k=1}^{N}\left|P_{meas,k} - P_{pred,k}\right|"),
        ("Root mean square error", "rmse", r"\mathrm{RMSE} = \sqrt{\frac{1}{N}\sum_{k=1}^{N}\left(P_{meas,k} - P_{pred,k}\right)^2}"),
        ("Mean absolute percentage error", "mape", r"\mathrm{MAPE} = \frac{100\%}{N}\sum_{k=1}^{N}\left|\frac{P_{meas,k} - P_{pred,k}}{P_{meas,k}}\right|"),
        ("Warning / abnormal limits", "thr", r"\theta_{warn} = \mu + 2\sigma, \qquad \theta_{abn} = \mu + 3\sigma"),
    ], "Used on healthy data now and on measured hardware data later.")

    # 03 HMI and dashboard
    d.section("HMI and Dashboard")
    d.diagram("Two UIs, one data bus", "d12-architecture",
              "Only the HMI sends commands. The dashboard is read-only.",
              "Later the ESP32 can publish to the same bus.")
    d.diagram("Screen map", "d13-screen-map", "", "Navigation across both UIs.")
    d.diagram("Setup flow", "d14-setup-flow",
              "Site, profile, parameters and sensors are all set in the UI.",
              "The user provides the site, input profile and plant parameters here.")
    d.wireframe("HMI: normal operation", "01-hmi-main", [
        "Plant and twin side by side",
        "Deviation panel between them",
        "Operator controls and fault buttons below",
        "State banner at the top",
    ], "Wireframe. Values are placeholders.")
    d.wireframe("HMI: model view", "05-hmi-model", [
        "SDM circuit, equation and parameters",
        "Live inputs and resulting curve",
        "Operating point moves with G and T",
    ], "The model is visible, not hidden in code.")
    d.wireframe("HMI: setup and site", "06-hmi-setup", [
        "Site: latitude, longitude, tilt, azimuth",
        "Input profile: preset, CSV, NASA POWER",
        "Datasheet, mismatch, sensors, thresholds",
    ], "Everything the run needs is set here.")
    d.wireframe("HMI: uploading a profile", "07-hmi-upload-dialog", [
        "CSV with timestamp, G, T",
        "Column mapping for other headers",
        "Automatic checks before use",
    ], "Checks catch bad files before the run starts.")
    d.wireframe("HMI: fault active", "02-hmi-fault-active", [
        "Banner turns ABNORMAL",
        "Twin still predicts the healthy value",
        "Alarm raised and logged",
    ], "This is the key moment of the demo.")
    d.wireframe("Dashboard: overview", "03-dashboard-overview", [
        "KPIs and power trends",
        "Deviation bands and I-V curve",
        "Time in each state and latest events",
    ], "Read-only view of the same data.")
    d.wireframe("Dashboard: validation", "04-dashboard-validation", [
        "MAE, RMSE and MAPE",
        "Error histogram and scatter",
        "Confusion matrix and detection delay",
    ], "Results come from the same run.")

    # 04 Live demo
    d.section("Live Demo")
    d.diagram("Demo 1: setup", "d20-seq-setup", "", "About 2 minutes.")
    d.diagram("Demo 2: normal operation", "d21-seq-normal", "Show the model tab and the dashboard overview here.", "About 2 minutes.")
    d.diagram("Demo 3: inject a fault", "d22-seq-fault", "Pause on the gap between the plant and the twin.", "About 3 minutes.")
    d.diagram("Demo 4: recovery", "d23-seq-recovery", "", "About 1 minute.")
    d.diagram("Demo 5: validation", "d24-seq-validation", "", "About 2 minutes.")

    d.diagram("Next steps", "d30-roadmap",
              "Simulation first, then hardware integration and experimental validation.",
              "Hardware stage uses MAE, RMSE and MAPE on real data.")
    d.finish()
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main(sys.argv[1])
